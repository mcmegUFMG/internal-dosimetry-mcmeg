"""
SAF Data Extractor
==================
Recursively scans data/phantoms/ directory to extract:
  Phantom Name -> Particle Name -> Source Organ -> Target Organ -> {
    energy: [...],
    phits: [...],
    rel_error: [...],
    sd: [...]
  }
Outputs: app/data.js and data.js
"""

import os
import sys
import glob
import re
import json
import math
import openpyxl

# Ensure utf-8 output in Windows terminals
if sys.stdout.encoding != 'utf-8':
    try:
        sys.stdout.reconfigure(encoding='utf-8')
    except Exception:
        pass

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DATA_DIR = os.path.join(BASE_DIR, "data", "phantoms")
APP_DATA_JS = os.path.join(BASE_DIR, "app", "data.js")
ROOT_DATA_JS = os.path.join(BASE_DIR, "data.js")


def clean_source_name(name):
    """Normalize sheet names to clean Source Organ names."""
    if not name:
        return None
    name = str(name).strip()
    if "setup" in name.lower() or "metadata" in name.lower():
        return None
    
    # Remove leading 'The '
    if name.lower().startswith("the "):
        name = name[4:]
        
    # Remove trailing ' as a Source Organ', ' as a S', ' as a Source', etc.
    name = re.sub(r"\s+as\s+a\s+source\s*organ\b.*$", "", name, flags=re.IGNORECASE)
    name = re.sub(r"\s+as\s+a\s+source\b.*$", "", name, flags=re.IGNORECASE)
    name = re.sub(r"\s+as\s+a\s+s\b.*$", "", name, flags=re.IGNORECASE)
    name = re.sub(r"\s+as\s+a\b.*$", "", name, flags=re.IGNORECASE)
    
    # Clean whitespace and capitalize properly
    parts = [p.capitalize() if not p.isupper() else p for p in name.split()]
    cleaned = " ".join(parts).strip()
    
    if cleaned.lower() == "urinary bladder wall":
        cleaned = "Urinary Bladder Wall"
    elif cleaned.lower() == "pituitary gland":
        cleaned = "Pituitary Gland"
        
    return cleaned if cleaned else None


def parse_float(val):
    """Parse numeric float or return None for N/A, blanks, NaN."""
    if val is None:
        return None
    if isinstance(val, (int, float)):
        if math.isnan(val) or math.isinf(val):
            return None
        return float(val)
    
    s = str(val).strip()
    if s == "" or s.lower() in ("n/a", "na", "null", "none", "-", "--"):
        return None
    
    s = s.replace(",", ".")
    try:
        f = float(s)
        if math.isnan(f) or math.isinf(f):
            return None
        return f
    except ValueError:
        return None


def extract_sheet_data(sheet):
    """Extract all target organ curves from a source organ sheet."""
    header_row = None
    for r in range(1, 15):
        c1 = sheet.cell(row=r, column=1).value
        if c1 and "target" in str(c1).lower():
            header_row = r
            break
            
    if not header_row:
        return {}

    source_data = {}
    current_target = None
    max_row = sheet.max_row

    for r in range(header_row + 1, max_row + 1):
        c1_val = sheet.cell(row=r, column=1).value
        if c1_val is not None and str(c1_val).strip() != "":
            raw_target = str(c1_val).strip()
            if "target" not in raw_target.lower():
                current_target = raw_target
                if current_target not in source_data:
                    source_data[current_target] = {
                        "energy": [],
                        "phits": [],
                        "rel_error": [],
                        "sd": []
                    }

        if current_target:
            e_val = parse_float(sheet.cell(row=r, column=2).value)
            if e_val is None:
                continue

            phits_val = parse_float(sheet.cell(row=r, column=3).value)
            rel_err_val = parse_float(sheet.cell(row=r, column=4).value)
            sd_val = parse_float(sheet.cell(row=r, column=7).value)

            # Auto calculate SD if missing in Col G but PHITS and Rel Error % exist
            if sd_val is None and phits_val is not None and rel_err_val is not None:
                sd_val = phits_val * (rel_err_val / 100.0)

            source_data[current_target]["energy"].append(e_val)
            source_data[current_target]["phits"].append(phits_val)
            source_data[current_target]["rel_error"].append(rel_err_val)
            source_data[current_target]["sd"].append(sd_val)

    # Clean out empty targets
    valid_targets = {}
    for tgt, d in source_data.items():
        if len(d["energy"]) > 0:
            valid_targets[tgt] = d

    return valid_targets


def extract_all():
    print(f"Scanning data directory: {DATA_DIR}")
    if not os.path.exists(DATA_DIR):
        print(f"Warning: {DATA_DIR} does not exist.")
        return {}

    database = {}

    # 1. Phantoms (Level 1 folders)
    phantom_dirs = [d for d in os.listdir(DATA_DIR) if os.path.isdir(os.path.join(DATA_DIR, d))]
    
    for p_name in sorted(phantom_dirs):
        p_path = os.path.join(DATA_DIR, p_name)
        database[p_name] = {}
        print(f"\n[Phantom] {p_name}")

        # 2. Particles (Level 2 folders)
        particle_dirs = [d for d in os.listdir(p_path) if os.path.isdir(os.path.join(p_path, d))]
        
        for part_name in sorted(particle_dirs):
            part_path = os.path.join(p_path, part_name)
            database[p_name][part_name] = {}
            print(f"  +-- [Particle] {part_name}")

            # 3. Excel Spreadsheets
            xlsx_files = glob.glob(os.path.join(part_path, "*.xlsx"))
            xlsx_files = [f for f in xlsx_files if not os.path.basename(f).startswith("~$")]

            for xlsx_file in xlsx_files:
                print(f"      +-- [File] {os.path.basename(xlsx_file)}")
                try:
                    wb = openpyxl.load_workbook(xlsx_file, data_only=True, read_only=False)
                    for sheet_name in wb.sheetnames:
                        src_name = clean_source_name(sheet_name)
                        if not src_name:
                            continue
                        
                        sheet = wb[sheet_name]
                        sheet_data = extract_sheet_data(sheet)
                        
                        if sheet_data:
                            database[p_name][part_name][src_name] = sheet_data
                            print(f"          |-- Source: {src_name:22} ({len(sheet_data)} targets)")
                except Exception as e:
                    print(f"          Error reading {xlsx_file}: {e}")

    # Output to JS files
    os.makedirs(os.path.dirname(APP_DATA_JS), exist_ok=True)
    
    js_content = "const INITIAL_DATA = " + json.dumps(database, ensure_ascii=False, indent=2) + ";\n"
    
    with open(APP_DATA_JS, "w", encoding="utf-8") as f:
        f.write(js_content)
    print(f"\n[Success] Wrote structured database to: {APP_DATA_JS}")

    with open(ROOT_DATA_JS, "w", encoding="utf-8") as f:
        f.write(js_content)
    print(f"[Success] Updated root data.js: {ROOT_DATA_JS}")

    return database


if __name__ == "__main__":
    extract_all()
